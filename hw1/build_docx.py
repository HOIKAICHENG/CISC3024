"""Build HW1_Solutions.docx with native Word equations and figures."""

from latex2mathml.converter import convert
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

XSL_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

_transform = etree.XSLT(etree.parse(XSL_PATH))


def latex_to_omath(latex: str):
    mathml = convert(latex)
    tree = etree.fromstring(mathml.encode("utf-8"))
    omml = _transform(tree)
    xml = etree.tostring(omml.getroot(), encoding="unicode")
    return parse_xml(xml)


def add_math(paragraph, latex: str):
    paragraph._p.append(latex_to_omath(latex))


def add_run(paragraph, text, *, bold=False, italic=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def mixed(doc, parts, *, space_after=6, space_before=0, bold=False, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    for part in parts:
        if isinstance(part, tuple) and part[0] == "m":
            add_math(p, part[1])
        else:
            add_run(p, part, bold=bold, size=size)
    return p


def display(doc, latex: str, space_before=8, space_after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    omath = latex_to_omath(latex)
    para = parse_xml(f'<m:oMathPara xmlns:m="{MATH_NS}"></m:oMathPara>')
    para.append(omath)
    p._p.append(para)
    return p


def set_cell_shading(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}" w:val="clear"/>')
    tcPr.append(shd)


def set_cell_margins(cell, **kw):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in kw.items():
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cell_clear(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return p


def cell_mixed(cell, parts, *, center=True, header=False):
    p = cell_clear(cell)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for part in parts:
        if isinstance(part, tuple) and part[0] == "m":
            add_math(p, part[1])
        else:
            add_run(p, part, bold=header, size=11)
    set_cell_margins(cell, top=40, bottom=40, left=60, right=60)


def format_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = parse_xml(
        f"""<w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>"""
    )
    tblPr.append(borders)
    if header:
        for cell in table.rows[0].cells:
            set_cell_shading(cell, "D9D9D9")
    for i, row in enumerate(table.rows):
        if i == 0:
            continue
        if i % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F2F2F2")


def fill_table(table, headers, rows):
    for j, h in enumerate(headers):
        cell_mixed(table.rows[0].cells[j], [h], header=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell_mixed(table.rows[i + 1].cells[j], [val])
    format_table(table)


def add_picture_centered(doc, path, width_in):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    add_run(p, text, italic=True, size=11)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Times New Roman"
    return h


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(fp, "Page ", size=11)
    add_page_number(fp)

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(2)
    add_run(t, "CISC3024 Pattern Recognition", bold=True, size=14)

    st = doc.add_paragraph()
    st.paragraph_format.space_after = Pt(8)
    add_run(st, "Homework 1", bold=True, size=14)

    info = doc.add_paragraph()
    info.paragraph_format.space_after = Pt(14)
    add_run(info, "Name: __________________     Student ID: __________________", size=12)

    # ----- 3.1 -----
    heading(doc, "3.1", 1)
    mixed(doc, [
        ("m", r"\omega_1=\{[0,0]^{T},[0,1]^{T}\}"),
        ",  ",
        ("m", r"\omega_2=\{[1,0]^{T},[1,1]^{T}\}"),
        ",  ",
        ("m", r"\rho=1"),
        ",  ",
        ("m", r"w(0)=[0,0]^{T}"),
        ".",
    ])
    mixed(doc, [
        "Decision hyperplane (lecture): ",
        ("m", r"g(x)=w^{T}x+w_0=0"),
        ".  ",
        ("m", r"w^{T}x+w_0>0 \Rightarrow x\in\omega_1"),
        ",  ",
        ("m", r"w^{T}x+w_0<0 \Rightarrow x\in\omega_2"),
        ".",
    ])
    mixed(doc, [
        "Use the same notation as the lecture example, ",
        ("m", r"w=[w_1,w_2,w_0]^{T}"),
        ". Augment each feature vector by 1:",
    ])
    display(doc, r"x=[x_1,x_2,1]^{T},\qquad w=[w_1,w_2,w_0]^{T}")
    mixed(doc, [
        "Then ", ("m", r"g(x)=w^{T}x"),
        ". Set ", ("m", r"w(0)=[0,0,0]^{T}"), ".",
    ])
    mixed(doc, ["Training vectors:"])
    display(doc, r"x^{(1)}=[0,0,1]^{T},\ x^{(2)}=[0,1,1]^{T}\in\omega_1")
    display(doc, r"x^{(3)}=[1,0,1]^{T},\ x^{(4)}=[1,1,1]^{T}\in\omega_2")
    mixed(doc, ["Reward-and-punishment form:"])
    mixed(doc, [
        "If ", ("m", r"x\in\omega_1"), " and ",
        ("m", r"w^{T}x\le 0"), ", then ",
        ("m", r"w\leftarrow w+\rho x"), ".",
    ])
    mixed(doc, [
        "If ", ("m", r"x\in\omega_2"), " and ",
        ("m", r"w^{T}x\ge 0"), ", then ",
        ("m", r"w\leftarrow w-\rho x"), ".",
    ])
    mixed(doc, ["Otherwise ", ("m", r"w"), " is unchanged."])
    mixed(doc, [
        "Present samples in the order ",
        ("m", r"x^{(1)},x^{(2)},x^{(3)},x^{(4)}"),
        ". Repeat until one full cycle has no update.",
    ])
    mixed(doc, [
        "t=1: ", ("m", r"x^{(1)}=[0,0,1]^{T}\in\omega_1"),
        ", ", ("m", r"w=[0,0,0]^{T}"),
        ", ", ("m", r"w^{T}x=0\le 0"),
        ". ", ("m", r"w=[0,0,0]^{T}+[0,0,1]^{T}=[0,0,1]^{T}"),
        ".",
    ])
    mixed(doc, [
        "t=2: ", ("m", r"x^{(2)}=[0,1,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=1>0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=3: ", ("m", r"x^{(3)}=[1,0,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=1\ge 0"),
        ". ", ("m", r"w=[0,0,1]^{T}-[1,0,1]^{T}=[-1,0,0]^{T}"),
        ".",
    ])
    mixed(doc, [
        "t=4: ", ("m", r"x^{(4)}=[1,1,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=-1<0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=5: ", ("m", r"x^{(1)}=[0,0,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=0\le 0"),
        ". ", ("m", r"w=[-1,0,0]^{T}+[0,0,1]^{T}=[-1,0,1]^{T}"),
        ".",
    ])
    mixed(doc, [
        "t=6: ", ("m", r"x^{(2)}=[0,1,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=1>0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=7: ", ("m", r"x^{(3)}=[1,0,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=0\ge 0"),
        ". ", ("m", r"w=[-1,0,1]^{T}-[1,0,1]^{T}=[-2,0,0]^{T}"),
        ".",
    ])
    mixed(doc, [
        "t=8: ", ("m", r"x^{(4)}=[1,1,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=-2<0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=9: ", ("m", r"x^{(1)}=[0,0,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=0\le 0"),
        ". ", ("m", r"w=[-2,0,0]^{T}+[0,0,1]^{T}=[-2,0,1]^{T}"),
        ".",
    ])
    mixed(doc, [
        "t=10: ", ("m", r"x^{(2)}=[0,1,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=1>0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=11: ", ("m", r"x^{(3)}=[1,0,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=-1<0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=12: ", ("m", r"x^{(4)}=[1,1,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=-1<0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=13: ", ("m", r"x^{(1)}=[0,0,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=1>0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=14: ", ("m", r"x^{(2)}=[0,1,1]^{T}\in\omega_1"),
        ", ", ("m", r"w^{T}x=1>0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=15: ", ("m", r"x^{(3)}=[1,0,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=-1<0"),
        ". No update.",
    ])
    mixed(doc, [
        "t=16: ", ("m", r"x^{(4)}=[1,1,1]^{T}\in\omega_2"),
        ", ", ("m", r"w^{T}x=-1<0"),
        ". No update.",
    ])
    mixed(doc, [
        "One full cycle has no update. The algorithm stops.",
    ])

    mixed(doc, [
        "The algorithm converges to ",
        ("m", r"w=[-2,0,1]^{T}"),
        ". The separating line is ",
        ("m", r"-2x_1+w_0=0"),
        ", i.e. ",
        ("m", r"-2x_1+1=0"),
        ", or ",
        ("m", r"x_1=\frac{1}{2}"),
        ".",
    ])
    mixed(doc, [
        "Check: for ", ("m", r"[0,0]^{T}"), " and ",
        ("m", r"[0,1]^{T}"), ", ",
        ("m", r"g=1>0"),
        " (", ("m", r"\omega_1"), "). For ",
        ("m", r"[1,0]^{T}"), " and ",
        ("m", r"[1,1]^{T}"), ", ",
        ("m", r"g=-1<0"),
        " (", ("m", r"\omega_2"), ").",
    ])
    add_picture_centered(doc, "fig_3_1.png", 4.3)
    caption(doc, "Figure 1. Decision line x1 = 1/2.")

    # ----- 3.2 -----
    heading(doc, "3.2", 1)
    mixed(doc, [
        "The three lines are realized by three perceptrons in the first layer. Activation function (lecture):",
    ])
    display(doc, r"f(v)=\begin{cases}1,& v\ge 0\\ 0,& v<0\end{cases}")
    display(doc, r"g_1(x)=x_1+x_2,\qquad y_1=f(g_1(x))")
    display(doc, r"g_2(x)=x_2-\frac{1}{4},\qquad y_2=f(g_2(x))")
    display(doc, r"g_3(x)=x_1-x_2,\qquad y_3=f(g_3(x))")
    mixed(doc, [
        "Synaptic weights of the first layer, written as ",
        ("m", r"[w_1,w_2,w_0]^{T}"), ":",
    ])
    display(doc, r"w_1=[1,1,0]^{T},\ w_2=[0,1,-1/4]^{T},\ w_3=[1,-1,0]^{T}")
    mixed(doc, [
        "Intersection points: ",
        ("m", r"(-1/4,1/4)"), ", ",
        ("m", r"(0,0)"), ", ",
        ("m", r"(1/4,1/4)"),
        ". The three lines form a triangle and divide the plane into 7 polyhedral regions. The pattern ",
        ("m", r"(g_1,g_2,g_3)=(-,+,+)"),
        " cannot occur: ",
        ("m", r"x_2>1/4"), " and ",
        ("m", r"x_1>x_2"), " imply ",
        ("m", r"x_1+x_2>1/2>0"), ".",
    ])
    mixed(doc, [
        "The first layer performs the mapping ",
        ("m", r"x\mapsto y=[y_1,y_2,y_3]^{T}"),
        ", ", ("m", r"y_i\in\{0,1\}"),
        ". Each region corresponds to one vertex of the unit cube.",
    ])
    mixed(doc, [
        "Example. Point ", ("m", r"(0,1/8)"),
        " (inside the triangle): ",
        ("m", r"g_1=1/8>0\Rightarrow y_1=1"),
        ", ",
        ("m", r"g_2=-1/8<0\Rightarrow y_2=0"),
        ", ",
        ("m", r"g_3=-1/8<0\Rightarrow y_3=0"),
        ". This region is mapped to ",
        ("m", r"(1,0,0)"), ".",
    ])
    mixed(doc, ["Mapping of all regions:"])

    t2 = doc.add_table(rows=8, cols=4)
    fill_table(
        t2,
        ["region", "test point", "(g1, g2, g3)", "cube vertex y"],
        [
            [("m", r"R_1"), "(0, 1/8)", ("m", r"(+,-,-)"), ("m", r"(1,0,0)")],
            [("m", r"R_2"), "(0, 1)", ("m", r"(+,+,-)"), ("m", r"(1,1,0)")],
            [("m", r"R_3"), "(-2, 1)", ("m", r"(-,+,-)"), ("m", r"(0,1,0)")],
            [("m", r"R_4"), "(2, 1)", ("m", r"(+,+,+)"), ("m", r"(1,1,1)")],
            [("m", r"R_5"), "(1, -0.2)", ("m", r"(+,-,+)"), ("m", r"(1,0,1)")],
            [("m", r"R_6"), "(-1, -0.2)", ("m", r"(-,-,-)"), ("m", r"(0,0,0)")],
            [("m", r"R_7"), "(0, -1)", ("m", r"(-,-,+)"), ("m", r"(0,0,1)")],
        ],
    )
    mixed(doc, [
        "Vertex ", ("m", r"(0,1,1)"),
        " is not used.",
    ], space_before=8)
    add_picture_centered(doc, "fig_3_2_regions.png", 5.5)
    caption(doc, "Figure 2. Three lines and seven polyhedral regions.")
    add_picture_centered(doc, "fig_3_2_cube.png", 5.2)
    caption(doc, "Figure 3. Mapping onto vertices of the unit cube.")
    mixed(doc, [
        "A two-layer perceptron: the output neuron realizes one hyperplane in the y-space and can separate some vertices from the others. A three-layer perceptron: the first layer forms hyperplanes, the second layer forms regions, the output neuron realizes an OR gate.",
    ])

    heading(doc, "(a) Two-layer network is sufficient", 2)
    mixed(doc, [
        "Let ",
        ("m", r"\omega_1=\{(0,0,0),(1,0,0),(0,1,0),(1,1,0)\}"),
        " and ",
        ("m", r"\omega_2=\{(0,0,1),(1,0,1),(1,1,1)\}"),
        ".",
    ])
    mixed(doc, [
        "These two sets of vertices are linearly separable by ",
        ("m", r"y_3=1/2"),
        ". The output perceptron is",
    ])
    display(doc, r"g(y)=-y_3+\frac{1}{2}=0")
    mixed(doc, [
        "If ", ("m", r"y_3=0"), ", then ",
        ("m", r"g=1/2>0"), " (", ("m", r"\omega_1"),
        "). If ", ("m", r"y_3=1"), ", then ",
        ("m", r"g=-1/2<0"), " (", ("m", r"\omega_2"), ").",
    ])
    mixed(doc, [
        "Output synaptic weights ",
        ("m", r"[w_{y1},w_{y2},w_{y3},w_0]^{T}=[0,0,-1,1/2]^{T}"),
        ". First-layer weights are given above.",
    ])

    heading(doc, "(b) Three-layer network is necessary", 2)
    mixed(doc, [
        "Let ",
        ("m", r"\omega_1=\{(0,0,0),(1,1,1)\}"),
        " and let ",
        ("m", r"\omega_2"),
        " be the remaining five realized vertices. This is the same situation as XOR: one hyperplane in y-space cannot put ",
        ("m", r"(0,0,0)"), " and ",
        ("m", r"(1,1,1)"),
        " on one side and the other vertices on the other side. Therefore a two-layer network is not sufficient.",
    ])
    mixed(doc, [
        "From the lecture: for each class-A vertex, construct a hyperplane that leaves that vertex on the (+) side and all others on the (-) side. The output neuron realizes an OR gate.",
    ])
    mixed(doc, [
        "Second layer, vertex ", ("m", r"(0,0,0)"), ":",
    ])
    display(doc, r"z_1=f(-y_1-y_2-y_3+\frac{1}{2}),\ w_{z1}=[-1,-1,-1,1/2]^{T}")
    mixed(doc, [
        "At ", ("m", r"(0,0,0)"), ", ",
        ("m", r"g=1/2>0"),
        ". At any other vertex, at least one ",
        ("m", r"y_i"), " equals 1, so ",
        ("m", r"g\le -1/2<0"), ".",
    ])
    mixed(doc, [
        "Second layer, vertex ", ("m", r"(1,1,1)"), ":",
    ])
    display(doc, r"z_2=f(y_1+y_2+y_3-\frac{5}{2}),\ w_{z2}=[1,1,1,-5/2]^{T}")
    mixed(doc, [
        "At ", ("m", r"(1,1,1)"), ", ",
        ("m", r"g=1/2>0"),
        ". At any other vertex, ",
        ("m", r"y_1+y_2+y_3\le 2"),
        ", so ",
        ("m", r"g\le -1/2<0"), ".",
    ])
    mixed(doc, [
        "Mapping: ",
        ("m", r"(0,0,0)\mapsto(z_1,z_2)=(1,0)"),
        ", ",
        ("m", r"(1,1,1)\mapsto(0,1)"),
        ", other vertices ",
        ("m", r"\mapsto(0,0)"),
        ". Output OR:",
    ])
    display(doc, r"g(z)=z_1+z_2-\frac{1}{2},\qquad w=[1,1,-1/2]^{T}")
    mixed(doc, [
        "If ", ("m", r"z_1"), " or ", ("m", r"z_2"),
        " equals 1, then ", ("m", r"g\ge 1/2>0"),
        " (", ("m", r"\omega_1"), ").",
    ])
    mixed(doc, ["Synaptic weights:"])

    t3 = doc.add_table(rows=7, cols=3)
    fill_table(
        t3,
        ["layer", "neuron", "w = [weights, w0]^T"],
        [
            ["1", ("m", r"y_1"), ("m", r"[1,1,0]^{T}")],
            ["1", ("m", r"y_2"), ("m", r"[0,1,-1/4]^{T}")],
            ["1", ("m", r"y_3"), ("m", r"[1,-1,0]^{T}")],
            ["2", ("m", r"z_1"), ("m", r"[-1,-1,-1,1/2]^{T}")],
            ["2", ("m", r"z_2"), ("m", r"[1,1,1,-5/2]^{T}")],
            ["3", "output (OR)", ("m", r"[1,1,-1/2]^{T}")],
        ],
    )

    out = "HW1_Solutions.docx"
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    build()
