import json
from pathlib import Path

import pytest
from docx import Document

import build_report
from build_report import SOURCE_URL, build_report as generate_report


ROOT = Path(__file__).resolve().parents[1]
FIGURE_NAMES = (
    "architecture.png",
    "training_curves.png",
    "reliability.png",
    "ood_entropy.png",
    "metrics_comparison.png",
    "uncertainty_examples.png",
)


def _document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join((*paragraphs, *cells))


@pytest.fixture(scope="module")
def quick_report():
    docx_path, pdf_path = generate_report(quick=True)
    document = Document(docx_path)
    return docx_path, pdf_path, document, _document_text(document)


def test_quick_report_has_required_sections_and_measured_results(quick_report):
    docx_path, pdf_path, _, text = quick_report
    for heading in (
        "1. How I Asked AI Tools to Find the Algorithm",
        "2. Algorithm Description",
        "3. How AI Implemented the Algorithm",
        "4. Experiment Settings and Results",
        "5. What I Have Learnt",
        "6. Webpage Link of the Source Codes",
    ):
        assert heading in text

    assert "HOI KAI CHENG" in text
    assert "UC325381" in text
    assert "SMOKE-TEST DATA" in text
    assert "NOT THE FINAL SUBMISSION" in text
    assert "CVPR 2024 paper findings" in text
    assert "Measured assignment results" in text

    evaluation = json.loads(
        (ROOT / "results" / "evaluation_quick.json").read_text(encoding="utf-8")
    )
    for method in ("map", "abnn_single", "abnn_full"):
        classification = evaluation["methods"][method]["classification"]
        for metric in ("accuracy", "nll", "brier", "ece"):
            assert f"{classification[metric]:.4f}" in text
        for dataset in ("svhn", "cifar100"):
            for metric in ("auroc", "aupr", "fpr95"):
                value = evaluation["methods"][method]["ood"][dataset][metric]
                assert f"{value:.4f}" in text

    assert docx_path.exists() and docx_path.stat().st_size > 100_000
    assert pdf_path.exists() and pdf_path.stat().st_size > 10_000


def test_quick_report_embeds_all_figures_tables_captions_and_live_code(
    quick_report,
):
    _, _, document, text = quick_report
    assert len(document.inline_shapes) == len(FIGURE_NAMES)
    assert len(document.tables) >= 5
    for number, figure_name in enumerate(FIGURE_NAMES, start=1):
        assert f"Figure {number}." in text
        assert figure_name.removesuffix(".png").replace("_", " ") in text.lower()
    assert "Table 1." in text
    assert "References" in text
    assert "Franchi" in text

    bnl_line = "scale = scale * (1.0 + torch.randn_like(scale))"
    averaging_line = "probabilities = sample_probs.mean(axis=0)"
    source = (ROOT / "code" / "abnn.py").read_text(encoding="utf-8")
    assert bnl_line in source and bnl_line in text
    assert averaging_line in source and averaging_line in text


def test_quick_report_contains_clickable_direct_source_link(quick_report):
    _, _, document, text = quick_report
    assert SOURCE_URL in text
    external_targets = {
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.is_external
    }
    assert SOURCE_URL in external_targets


def test_word_conversion_closes_document_and_application_on_error(
    tmp_path, monkeypatch
):
    source = tmp_path / "source.docx"
    source.write_bytes(b"placeholder")
    target = tmp_path / "target.pdf"

    class FakeDocument:
        def __init__(self):
            self.closed = False
            self.Fields = type("Fields", (), {"Update": lambda self: None})()

        def ExportAsFixedFormat(self, *_args, **_kwargs):
            raise OSError("simulated Word export failure")

        def Close(self, _save_changes):
            self.closed = True

    class FakeWord:
        def __init__(self):
            self.document = FakeDocument()
            self.Documents = type(
                "Documents",
                (),
                {"Open": lambda _self, _path: self.document},
            )()
            self.Visible = True
            self.DisplayAlerts = 1
            self.quit_called = False

        def Quit(self):
            self.quit_called = True

    word = FakeWord()
    monkeypatch.setattr(build_report, "_dispatch_word", lambda: word)

    with pytest.raises(RuntimeError, match="Microsoft Word.*source.docx"):
        build_report._convert_docx_to_pdf(source, target)

    assert word.document.closed
    assert word.quit_called


def test_word_conversion_uses_fixed_format_export_and_cleans_up(
    tmp_path, monkeypatch
):
    source = tmp_path / "source.docx"
    source.write_bytes(b"placeholder")
    target = tmp_path / "target.pdf"

    class FakeDocument:
        def __init__(self):
            self.closed = False
            self.fields_updated = False
            self.Fields = type(
                "Fields",
                (),
                {"Update": lambda _self: setattr(self, "fields_updated", True)},
            )()

        def ExportAsFixedFormat(self, output_path, export_format):
            assert export_format == 17
            Path(output_path).write_bytes(b"%PDF-1.7\nreport")

        def SaveAs2(self, *_args, **_kwargs):
            raise AssertionError("SaveAs2 is less reliable for PDF export")

        def Close(self, _save_changes):
            self.closed = True

    class FakeWord:
        def __init__(self):
            self.document = FakeDocument()
            self.Documents = type(
                "Documents",
                (),
                {"Open": lambda _self, _path: self.document},
            )()
            self.quit_called = False

        def Quit(self):
            self.quit_called = True

    word = FakeWord()
    monkeypatch.setattr(build_report, "_dispatch_word", lambda: word)

    assert build_report._convert_docx_to_pdf(source, target) == target.resolve()
    assert word.document.fields_updated
    assert word.document.closed
    assert word.quit_called
